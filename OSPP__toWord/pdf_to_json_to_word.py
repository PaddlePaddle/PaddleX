# Copyright (c) 2024 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from typing import List, Dict
import json, os, re, copy

TRANSLATABLE_LABELS = {"chart", "image", "seal", "number"}
SPLIT_TOKEN = "¥$¥"


# --- 样式设置 ---
# 设置段落的字体、字号、加粗、对齐方式和首行缩进
def set_paragraph_style(para, font_name="Times New Roman", font_size_pt=12, bold=False, indent=False, alignment = None):
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    run = para.runs[0] if para.runs else para.add_run()
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(font_size_pt)
    run.bold = bold
    if alignment is None:
        alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.alignment = alignment
    if indent:
        para.paragraph_format.first_line_indent = Inches(0.3)

# 设置 run 的字体，包括中英文字体和字号
def set_run_font(run, font_name_en="Times New Roman", font_name_cn="宋体", font_size_pt=10.5, bold=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    
    run.font.name = font_name_en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name_cn)
    run.font.size = Pt(font_size_pt)
    run.bold = bold

# 清空并设置 section 的页眉或页脚内容，居中显示
def set_section_part_text(section_part, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for _ in range(len(section_part.paragraphs)):
        p = section_part.paragraphs[0]
        p._element.getparent().remove(p._element)
    para = section_part.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run_font(run)


# --- 内容格式 ---
# 根据块的标签和内容，向文档添加对应格式的段落或标题
def format_block_style(doc, label, content):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    style_map = {
        "doc_title":     {"level": 0, "size": 20, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "header":        {"size": 16, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "abstract_title": {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "content_title":  {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "reference_title": {"level": 1, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "paragraph_title": {"level": 2, "size": 14, "bold": True, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "abstract":      {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
        "text":          {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY, "indent": True},
        "figure_title":  {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "table_title":   {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "chart_title":   {"size": 10, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "reference":     {"size": 12, "align": WD_ALIGN_PARAGRAPH.JUSTIFY},
        "algorithm":     {"font": "Courier New", "size": 11, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "formula":       {"size": 12, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "vision_footnote": {"size": 9, "align": WD_ALIGN_PARAGRAPH.LEFT},
        "number":        {"size": 9, "align": WD_ALIGN_PARAGRAPH.CENTER},
        "footer":        {"size": 9, "align": WD_ALIGN_PARAGRAPH.CENTER}
    }

    config = style_map.get(label, {"size": 12, "align": WD_ALIGN_PARAGRAPH.LEFT})
    para = doc.add_heading(content, level=config["level"]) if "level" in config else doc.add_paragraph(content)
    set_paragraph_style(para,
                        font_name=config.get("font", "Times New Roman"),
                        font_size_pt=config["size"],
                        bold=config.get("bold", False),
                        indent=config.get("indent", False),
                        alignment=config["align"])


# --- 表格解析 ---
# 解析 HTML 表格字符串，提取为二维文本列表
def parse_html_table(html):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    return [[cell.get_text(strip=True) for cell in tr.find_all(["td", "th"])]
            for tr in soup.find_all("tr")]


# --- 排序 ---
# 根据页码和块的纵坐标，对块进行排序
def sort_blocks_by_position(blocks):
    return sorted(blocks, key=lambda b: (b.get("page", 0), b["block_bbox"][1]))

# 从 markdown 文件中提取指定图片的宽度比例，默认 1.0
def get_image_width_from_md(md_path, image_name):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    pattern = re.compile(rf'<img\s+[^>]*src=["\'].*?{re.escape(image_name)}.*?["\'][^>]*width=["\'](\d+)%["\']', re.I)
    match = pattern.search(content)
    return int(match.group(1)) / 100 if match else 1.0

# 向文档插入图片，宽度根据比例缩放，居中显示
def insert_image(doc, image_path, width_ratio):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_picture(image_path, width=Inches(5.5 * width_ratio))
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


# --- 主流程函数 ---
# 从 JSON 读取块列表，按页生成 Word 文档，支持页眉页脚、表格、图片等
def blocks_to_word(json_path, word_output_path, image_base_path, input_path, output_path):
    from docx import Document
    from docx.enum.section import WD_SECTION
    
    with open(json_path, "r", encoding="utf-8") as f:
        blocks = json.load(f)

    doc = Document()
    pages = {}
    for block in blocks:
        pages.setdefault(block.get("page", 0), []).append(block)

    for page_num in sorted(pages.keys()):
        page_blocks = sort_blocks_by_position(pages[page_num])
        section = doc.add_section(WD_SECTION.NEW_PAGE) if page_num != 0 else doc.sections[0]

        # 页眉
        header_blocks = [b for b in page_blocks if b["block_label"] == "header"]
        if header_blocks:
            header_text = "\n".join(b["block_content"].strip() for b in header_blocks if b.get("block_content"))
            section.header.is_linked_to_previous = False
            set_section_part_text(section.header, header_text)

        # 页脚
        footer_blocks = [b for b in page_blocks if b["block_label"] == "footer"]
        if footer_blocks:
            footer_text = "\n".join(b["block_content"].strip() for b in footer_blocks if b.get("block_content"))
            section.footer.is_linked_to_previous = False
            set_section_part_text(section.footer, footer_text)

        for block in page_blocks:
            label = block["block_label"]
            content = block.get("block_content", "").strip()
            if not content and label not in ["chart", "image", "table", "seal"]:
                continue

            if label in ["chart", "image", "seal"]:
                bbox = block.get("block_bbox", [0, 0, 0, 0])
                x1, y1, x2, y2 = map(int, bbox)
                filename1 = f"img_in_chart_box_{x1}_{y1}_{x2}_{y2}.jpg"
                filename2 = f"img_in_image_box_{x1}_{y1}_{x2}_{y2}.jpg"
                image_filename = filename1 if os.path.exists(os.path.join(image_base_path, filename1)) else filename2
                image_path = os.path.join(image_base_path, image_filename)

                if os.path.exists(image_path):
                    base_name = os.path.splitext(os.path.basename(input_path))[0]
                    md_path = f"{output_path}/{base_name}_{block.get('page')}.md"
                    width = get_image_width_from_md(md_path, image_filename)
                    insert_image(doc, image_path, width)
                else:
                    doc.add_paragraph(f"[Image {image_filename} not found]")
                continue

            elif label == "table":
                rows = parse_html_table(content) if "<table" in content else [
                    row.split('\t') for row in content.strip().split('\n') if row.strip()
                ]
                if rows:
                    max_cols = max(len(r) for r in rows)  # 考虑每行列数不一样的情况
                    table = doc.add_table(rows=0, cols=max_cols)
                    table.style = "Table Grid"
                    for row_cells in rows:
                        row = table.add_row().cells
                        for i in range(max_cols):
                            if i < len(row_cells):  # 正常填充
                                row[i].text = row_cells[i].strip()
                            else:  # 如果该行缺列，用空字符串补齐
                                row[i].text = ""
                continue

            if label not in ["header", "footer"]:
                format_block_style(doc, label, content)

    doc.save(word_output_path)
    print(f"保存 Word 至：{word_output_path}")


# --- JSON 合并函数 ---
# 将多个 JSON 结果合并为一个，添加 page 字段，保存合并文件并返回路径
def merge_block(json_list: List[Dict], output_path) -> str:
    merged = []
    for i, item in enumerate(json_list):
        blocks = copy.deepcopy(item["res"]["parsing_res_list"])
        for b in blocks:
            b["page"] = i
        merged.extend(blocks)

    merged_path = os.path.join(output_path, "merged_translated.json")
    with open(merged_path, "w", encoding="utf-8") as f:
        json.dump(merged, f, ensure_ascii=False, indent=4)
    return merged_path

def md2cn_word(json_list: List[Dict],input_path, output_path):
    # 将所有的json中的block抽取出来，合并成一个json
    merged_json_path = merge_block(json_list,output_path=output_path)

    # 调用转换成 Word 的函数
    base_name = os.path.splitext(os.path.basename(input_path))[0]

    blocks_to_word(
        json_path=merged_json_path,
        word_output_path=f"{output_path}/{base_name}_toword.docx",
        image_base_path=f"{output_path}/imgs",
        input_path=input_path,
        output_path=output_path,
    )