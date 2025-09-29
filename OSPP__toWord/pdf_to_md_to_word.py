# Copyright (c) 2024 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import os, re


def set_paragraph_style(paragraph, bold=False, align="left", font_size=11, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.shared import RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    """统一设置段落样式"""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")  # 中文宋体
    run.font.size = Pt(font_size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align == "center":
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    elif align == "right":
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    else:
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT


def add_image(paragraph, src, width_percent):
    from docx.shared import Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    """插入图片并缩放"""
    if os.path.exists(src):
        try:
            width_in_inches = Inches(width_percent / 100 * 6.0)
            run = paragraph.add_run()
            run.add_picture(src, width=width_in_inches)
            paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        except Exception as e:
            paragraph.add_run(f"[图片加载失败: {src}]")
    else:
        paragraph.add_run(f"[图片不存在: {src}]")


def add_table(document, table_html):
    from bs4 import BeautifulSoup

    """解析 HTML 表格并添加到 Word"""
    soup = BeautifulSoup(table_html, "html.parser")
    table_tag = soup.find("table")
    if not table_tag:
        return

    rows = table_tag.find_all("tr")
    if not rows:
        return

    # 计算最大列数，保证不会越界
    max_cols = max(len(row.find_all(["td", "th"])) for row in rows)
    table = document.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"

    for i, row in enumerate(rows):
        cells = row.find_all(["td", "th"])
        for j in range(max_cols):
            if j < len(cells):
                text = cells[j].get_text(strip=True)
                table.cell(i, j).text = text
            else:
                table.cell(i, j).text = ""  # 列数不足补空


from typing import Dict


def json_to_html_with_headfoot(json_list: list, input_path: str) -> Dict:

    # HTML 表格头
    html_lines = [
        "<html><body><table border='1'>",
        "<tr><th>page</th><th>header</th><th>footer</th><th>footnote</th><th>page_number</th></tr>",
    ]

    for page_idx, json_blocks in enumerate(json_list, start=1):

        parsing_res_list = json_blocks.get("res", {}).get("parsing_res_list", [])

        head_foot_dict = {"header": "", "footer": "", "footnote": "", "page_number": ""}

        for block in parsing_res_list:
            label = block.get("block_label", "").lower()
            content = block.get("block_content", "").strip()
            if not content:
                continue

            if label in {"header", "footer", "footnote", "number", "page_number"}:
                if label == "number":
                    label = "page_number"
                head_foot_dict[label] = content

        # 构造表格行
        html_lines.append(
            f"<tr><td>{page_idx}</td>"
            f"<td>{head_foot_dict['header']}</td>"
            f"<td>{head_foot_dict['footer']}</td>"
            f"<td>{head_foot_dict['footnote']}</td>"
            f"<td>{head_foot_dict['page_number']}</td></tr>"
        )

    html_lines.append("</table></body></html>")

    # 将 HTML 行列表拼接成单个字符串
    html_string = "".join(html_lines)

    result = {
        "markdown_images": {},
        "page_index": 0,
        "input_path": input_path,
        "markdown_texts": html_string,  # 单个字符串
        "page_continuation_flags": (True, True),  # 可根据需要调整
    }
    return result


def process_md_page(document, md_text, output_path):
    from bs4 import BeautifulSoup

    """处理单页内容"""
    lines = md_text.strip().split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            continue

        title_color = (0, 0, 255)
        if line.startswith("##### "):
            p = document.add_paragraph(line[6:])
            set_paragraph_style(p, bold=True, font_size=10)
        elif line.startswith("#### "):
            p = document.add_paragraph(line[5:])
            set_paragraph_style(p, bold=True, font_size=11)
        elif line.startswith("### "):
            p = document.add_paragraph(line[4:])
            set_paragraph_style(p, bold=True, font_size=12)
        elif line.startswith("## "):
            p = document.add_paragraph(line[3:])
            set_paragraph_style(p, bold=True, font_size=14)
        elif line.startswith("# "):
            p = document.add_paragraph(line[2:])
            set_paragraph_style(p, bold=True, font_size=16)

        # 居中内容处理
        elif line.startswith("<div") and "text-align: center" in line:
            soup = BeautifulSoup(line, "html.parser")
            div = soup.find("div")
            if not div:
                continue
            if div.img:
                img = div.img
                src = img.get("src")
                width_attr = img.get("width", "100%").replace("%", "")
                width_percent = float(width_attr) if width_attr else 100
                p = document.add_paragraph()
                add_image(p, f"{output_path}/{src}", width_percent)
            elif div.table:
                add_table(document, str(div))
            else:
                text = div.get_text(strip=True)
                if text:
                    p = document.add_paragraph(text)
                    set_paragraph_style(p, bold=True, align="center", color=title_color)

        # HTML表格处理
        elif "<table" in line:
            add_table(document, line)

        # 普通段落
        else:
            p = document.add_paragraph(line)
            set_paragraph_style(p, font_size=11)


def md_to_word(md_text, output_path, base_name):
    from bs4 import BeautifulSoup
    from docx import Document

    pages = [p.strip() for p in re.split(r"<sep>1</sep>", md_text) if p.strip()]

    if not pages:
        print("❌ 没有有效内容")
        return

    # 最后一页如果包含表格，作为页眉页脚信息
    meta_info = pages[-1] if "<table" in pages[-1] else ""
    # 根据是否有页眉页脚，决定内容是否包含最后一段。有的话，正文内容，不包含最后一段
    content_pages = pages[:-1] if meta_info else pages

    if len(pages) == 1:
        content_pages = pages
        meta_info = None

    # 页眉页脚存在，则提取出来
    headers, footers, footnotes, page_numbers = [], [], [], []
    if meta_info:
        soup = BeautifulSoup(meta_info, "html.parser")
        table = soup.find("table")
        if table:
            trs = table.find_all("tr")
            for tr in trs[1:]:
                tds = tr.find_all("td")
                header_text = tds[1].get_text(strip=True)
                footer_text = tds[2].get_text(strip=True)
                footnote_text = tds[3].get_text(strip=True)
                page_number = tds[4].get_text(strip=True)
                headers.append(header_text)
                footers.append(footer_text)
                footnotes.append(footnote_text)
                page_numbers.append(page_number)

    document = Document()

    # 逐页处理
    for idx, page in enumerate(content_pages):

        section = document.sections[0] if idx == 0 else document.add_section()
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

        # 页眉
        if headers and idx < len(headers) and headers[idx]:
            header_para = (
                section.header.paragraphs[0]
                if section.header.paragraphs
                else section.header.add_paragraph()
            )
            header_para.text = headers[idx]
            set_paragraph_style(header_para, align="center")

        # 页脚
        if footers and idx < len(footers) and footers[idx]:
            footer_para = (
                section.footer.paragraphs[0]
                if section.footer.paragraphs
                else section.footer.add_paragraph()
            )
            footer_para.text = footers[idx]
            set_paragraph_style(footer_para, align="center")

        # 内容处理
        process_md_page(document, page, output_path)

        # 页脚部分：footnote + page number，作为正文插入
        if footnotes and idx < len(footnotes) and footnotes[idx]:
            footnote_para = document.add_paragraph(f"[footnote] {footnotes[idx]}")
            set_paragraph_style(footnote_para, align="left", font_size=9)

        if page_numbers and idx < len(page_numbers) and page_numbers[idx]:
            page_num_para = document.add_paragraph(page_numbers[idx])
            set_paragraph_style(page_num_para, align="center", font_size=9)

        if idx != len(content_pages) - 1:
            document.add_page_break()

    output_path = f"{output_path}/{base_name}_word.docx"
    document.save(output_path)
    print(f"✅ Word 已生成: {output_path}")
