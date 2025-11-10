# Copyright (c) 2024 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

import os
import re
from pathlib import Path

from ...common.result import BaseCVResult, LatexMixin, MarkdownMixin, WordMixin


class MarkdownResult(BaseCVResult, MarkdownMixin, WordMixin, LatexMixin):
    def __init__(self, data) -> None:
        """Initializes a new instance of the class with the specified data."""
        super().__init__(data)
        MarkdownMixin.__init__(self)

    def _get_input_fn(self):
        fn = super()._get_input_fn()
        if (page_idx := self.get("page_index", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{page_idx}{suffix}"
        if (language := self.get("language", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{language}{suffix}"
        return fn

    def _to_markdown(self, pretty=True, show_formula_number=False) -> dict:
        return self

    # in order to make MarkdownResult support save_to_word
    def _to_word(self, save_path) -> dict:
        from bs4 import BeautifulSoup

        md_text = self.get("markdown_texts", "")

        def set_paragraph_style(
            paragraph, bold=False, align="left", font_size=11, color=None
        ):
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
            from docx.oxml.ns import qn
            from docx.shared import Pt, RGBColor

            # Set paragraph style uniformly
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(font_size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor(*color)
            if align == "center":
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            elif align == "right":
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
            else:
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        def add_image(paragraph, src, width_percent):
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
            from docx.shared import Inches

            if os.path.exists(src):
                try:
                    width_in_inches = Inches(width_percent / 100 * 6.0)
                    run = paragraph.add_run()
                    run.add_picture(src, width=width_in_inches)
                    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                except Exception as e:
                    paragraph.add_run(f"[fail load image: {src}]")
            else:
                paragraph.add_run(f"[image not exist: {src}]")

        def add_table(document, table_html):
            """
            Parsing HTML table and add to Word
            """
            soup = BeautifulSoup(table_html, "html.parser")
            table_tag = soup.find("table")
            if not table_tag:
                return

            rows = table_tag.find_all("tr")
            if not rows:
                return

            # Calculate the maximum number of columns to avoid out-of-bounds errors
            max_cols = max(len(row.find_all(["td", "th"])) for row in rows)
            table = document.add_table(rows=len(rows), cols=max_cols)
            table.style = "Table Grid"

            for i, row in enumerate(rows):
                cells = row.find_all(["td", "th"])
                for j in range(max_cols):
                    if j < len(cells):
                        text = cells[j].get_text(strip=True)
                        table.cell(i, j).text = text
                    else:
                        table.cell(i, j).text = ""

        def process_md_page(document, md_text, output_path):

            # Process single page conten
            lines = md_text.strip().split("\n")
            for line in lines:
                line = line.strip()
                if not line:
                    continue

                title_color = (0, 0, 255)
                if line.startswith("##### "):
                    p = document.add_paragraph(line[6:])
                    set_paragraph_style(p, bold=True, font_size=10)
                elif line.startswith("#### "):
                    p = document.add_paragraph(line[5:])
                    set_paragraph_style(p, bold=True, font_size=11)
                elif line.startswith("### "):
                    p = document.add_paragraph(line[4:])
                    set_paragraph_style(p, bold=True, font_size=12)
                elif line.startswith("## "):
                    p = document.add_paragraph(line[3:])
                    set_paragraph_style(p, bold=True, font_size=14)
                elif line.startswith("# "):
                    p = document.add_paragraph(line[2:])
                    set_paragraph_style(p, bold=True, font_size=16)

                # Handle centered content
                elif line.startswith("<div") and "text-align: center" in line:
                    soup = BeautifulSoup(line, "html.parser")
                    div = soup.find("div")
                    if not div:
                        continue
                    if div.img:
                        img = div.img
                        src = img.get("src")
                        width_attr = img.get("width", "100%").replace("%", "")
                        width_percent = float(width_attr) if width_attr else 100
                        p = document.add_paragraph()
                        add_image(p, f"{output_path}/{src}", width_percent)
                    elif div.table:
                        add_table(document, str(div))
                    else:
                        text = div.get_text(strip=True)
                        if text:
                            p = document.add_paragraph(text)
                            set_paragraph_style(
                                p, bold=True, align="center", color=title_color
                            )

                # Handle HTML tables
                elif "<table" in line:
                    add_table(document, line)

                # Normal paragraph
                else:
                    p = document.add_paragraph(line)
                    set_paragraph_style(p, font_size=11)

        from docx import Document

        document = Document()
        process_md_page(document, md_text, save_path)

        return document

    # in order to make MarkdownResult support save_to_latex
    def _to_latex(self, save_path) -> str:
        from bs4 import BeautifulSoup

        def escape_latex_outside_formula(s: str) -> str:
            """
            Escape LaTeX special characters while preserving formulas.
            """
            if not s:
                return ""

            placeholders = []

            def repl(m):
                placeholders.append(m.group(0))
                return f"@@FORMULA{len(placeholders)-1}@@"

            # Extract formulas
            formula_pat = re.compile(
                r"(\$\$.*?\$\$|\$.*?\$|\\\[.*?\\\]|\\\(.*?\\\))", re.DOTALL
            )
            tmp = formula_pat.sub(repl, s)

            tmp = (
                tmp.replace("\\", "\\textbackslash{}")
                .replace("&", "\\&")
                .replace("%", "\\%")
                .replace("$", "\\$")
                .replace("#", "\\#")
                .replace("_", "\\_")
                .replace("{", "\\{")
                .replace("}", "\\}")
                .replace("~", "\\textasciitilde{}")
                .replace("^", "\\textasciicircum{}")
            )

            # Restore formulas
            for i, f in enumerate(placeholders):
                tmp = tmp.replace(f"@@FORMULA{i}@@", f)
            return tmp

        #
        def get_image_width_from_md_line(line, default_ratio=0.8):
            """
            Parse the image width attribute.
            """
            m = re.search(r'width\s*=\s*["\']?(\d+)%?["\']?', line)
            if m:
                val = int(m.group(1))
                return max(0.01, min(val / 100.0, 1.0))
            m2 = re.search(r"width\s*:\s*(\d+)%", line)
            if m2:
                val = int(m2.group(1))
                return max(0.01, min(val / 100.0, 1.0))
            return default_ratio

        def process_table_html(content) -> str:
            """
            Process table content.
            """
            if "<table" in content:
                soup = BeautifulSoup(content, "html.parser")
                rows = []
                for tr in soup.find_all("tr"):
                    row = []
                    for td in tr.find_all(["td", "th"]):
                        text = td.get_text(strip=True)
                        row.append(escape_latex_outside_formula(text))
                    rows.append(row)
            else:
                rows = [
                    [escape_latex_outside_formula(c) for c in row.split("\t")]
                    for row in content.splitlines()
                    if row.strip()
                ]

            if not rows:
                return ""

            col_count = max(len(r) for r in rows)
            norm_rows = [r + [""] * (col_count - len(r)) for r in rows]
            col_format = " ".join(
                [">{\\raggedright\\arraybackslash}X" for _ in range(col_count)]
            )

            latex = "\\begin{center}\n\\renewcommand{\\arraystretch}{1.5}\n"
            latex += f"\\begin{{tabularx}}{{\\textwidth}}{{{col_format}}}\n\\toprule\n"
            for i, row in enumerate(norm_rows):
                latex += " & ".join(row) + " \\\\\n"
                if i == 0:
                    latex += "\\midrule\n"
            latex += "\\bottomrule\n\\end{tabularx}\n\\end{center}\n\n"
            return latex

        def process_paragraph(s: str) -> str:
            """
            Process text paragraphs, preserving formulas.
            """
            paragraphs = re.split(r"\n\s*\n", s)
            processed_paras = []
            for p in paragraphs:
                p = p.strip()
                if not p:
                    continue
                processed_paras.append("\\par " + escape_latex_outside_formula(p))
            return "\n\n".join(processed_paras) + "\n\n"

        def process_md_line(line: str, save_path) -> str:
            """
            Process a single line.
            """
            line = line.strip()
            if not line:
                return ""

            if line.startswith("##### "):
                return f"\\paragraph*{{{escape_latex_outside_formula(line[6:].strip())}}}\n\n"
            if line.startswith("#### "):
                return f"\\subsubsection*{{{escape_latex_outside_formula(line[5:].strip())}}}\n\n"
            if line.startswith("### "):
                return f"\\subsection*{{{escape_latex_outside_formula(line[4:].strip())}}}\n\n"
            if line.startswith("## "):
                return f"\\section*{{{escape_latex_outside_formula(line[3:].strip())}}}\n\n"
            if line.startswith("# "):
                return f"\\section*{{{escape_latex_outside_formula(line[2:].strip())}}}\n\n"

            if "<div" in line and "text-align: center" in line:
                soup = BeautifulSoup(line, "html.parser")
                div = soup.find("div")
                if div:
                    if div.img:
                        img = div.img
                        src = img.get("src")
                        src = f"{save_path}/{src}"
                        width_ratio = get_image_width_from_md_line(str(img))
                        return (
                            f"\\begin{{figure}}[h]\n\\centering\n"
                            f"\\includegraphics[width={width_ratio:.2f}\\linewidth]{{{src}}}\n"
                            f"\\end{{figure}}\n\n"
                        )
                    if div.table:
                        return process_table_html(str(div))
                    text = div.get_text(strip=True)
                    if text:
                        return f"\\begin{{center}}{escape_latex_outside_formula(text)}\\end{{center}}\n\n"

            if "<table" in line:
                return process_table_html(line)
            return process_paragraph(line)

        latex_lines = [
            "\\documentclass[12pt]{article}",
            "\\usepackage{xeCJK}",
            "\\usepackage{fontspec}",
            "\\usepackage{graphicx}",
            "\\usepackage{amsmath}",
            "\\usepackage{geometry}",
            "\\usepackage{fancyhdr}",
            "\\usepackage{indentfirst}",
            "\\usepackage{caption}",
            "\\usepackage{tabularx, booktabs}",
            "\\usepackage{amssymb}",
            "\\usepackage{amsfonts}",
            "\\geometry{a4paper, margin=1in}",
            "\\setCJKmainfont{Droid Sans Fallback}",
            "\\setmainfont{DejaVu Serif}",
            "\\setsansfont{Lato}",
            "\\setmonofont{Latin Modern Mono}",
            "\\pagestyle{fancy}",
            "\\setlength{\\parindent}{2em}",
            "\\begin{document}\n",
        ]

        md_text = self.get("markdown_texts", "")
        for line in md_text.splitlines():
            latex_lines.append(process_md_line(line, save_path))

        latex_lines.append("\\end{document}")

        return "\n".join(latex_lines)


class DocumentResult(BaseCVResult, WordMixin):
    def __init__(self, data) -> None:
        """
        Initializes a new instance of the class with the specified data.
        """
        super().__init__(data)
        WordMixin.__init__(self)

    def _get_input_fn(self):
        fn = super()._get_input_fn()
        if (page_idx := self.get("page_index", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{page_idx}{suffix}"
        if (language := self.get("language", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{language}{suffix}"
        return fn

    def _to_word(self) -> dict:
        return self


class LatexResult(BaseCVResult, LatexMixin):
    def __init__(self, data) -> None:
        super().__init__(data)
        LatexMixin.__init__(self)

    def _get_input_fn(self):
        fn = super()._get_input_fn()
        if (page_idx := self.get("page_index", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{page_idx}{suffix}"
        if (language := self.get("language", None)) is not None:
            fp = Path(fn)
            stem, suffix = fp.stem, fp.suffix
            fn = f"{stem}_{language}{suffix}"
        return fn

    def _to_latex(self) -> dict:
        return self
