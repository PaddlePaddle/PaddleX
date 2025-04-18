# Copyright (c) 2024 PaddlePaddle Authors. All Rights Reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

from __future__ import absolute_import

from ....utils.deps import function_requires_deps, is_dep_available
from .style import Table

if is_dep_available("lxml"):
    from lxml import html
if is_dep_available("openpyxl"):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
if is_dep_available("premailer"):
    from premailer import Premailer


def string_to_int(s):
    """
    Convert a string to an integer
    """
    if s.isdigit():
        return int(s)
    return 0


@function_requires_deps("lxml")
def get_Tables(doc):
    """
    Find all the tables in the doc
    """
    tree = html.fromstring(doc)
    comments = tree.xpath("//comment()")
    for comment in comments:
        comment.drop_tag()
    return [Table(table) for table in tree.xpath("//table")]


@function_requires_deps("openpyxl")
def write_rows(worksheet, elem, row, column=1):
    """
    Writes every tr child element of elem to a row in the worksheet
    returns the next row after all rows are written
    """
    from openpyxl.cell.cell import MergedCell

    initial_column = column
    for table_row in elem.rows:
        for table_cell in table_row.cells:
            cell = worksheet.cell(row=row, column=column)
            while isinstance(cell, MergedCell):
                column += 1
                cell = worksheet.cell(row=row, column=column)

            colspan = string_to_int(table_cell.element.get("colspan", "1"))
            rowspan = string_to_int(table_cell.element.get("rowspan", "1"))
            if rowspan > 1 or colspan > 1:
                worksheet.merge_cells(
                    start_row=row,
                    start_column=column,
                    end_row=row + rowspan - 1,
                    end_column=column + colspan - 1,
                )

            cell.value = table_cell.value
            table_cell.format(cell)
            min_width = table_cell.get_dimension("min-width")
            max_width = table_cell.get_dimension("max-width")

            if colspan == 1:
                # Initially, when iterating for the first time through the loop, the width of all the cells is None.
                # As we start filling in contents, the initial width of the cell (which can be retrieved by:
                # worksheet.column_dimensions[get_column_letter(column)].width) is equal to the width of the previous
                # cell in the same column (i.e. width of A2 = width of A1)
                width = max(
                    worksheet.column_dimensions[get_column_letter(column)].width or 0,
                    len(table_cell.value) + 2,
                )
                if max_width and width > max_width:
                    width = max_width
                elif min_width and width < min_width:
                    width = min_width
                worksheet.column_dimensions[get_column_letter(column)].width = width
            column += colspan
        row += 1
        column = initial_column
    return row


def table_to_sheet(table, wb):
    """
    Takes a table and workbook and writes the table to a new sheet.
    The sheet title will be the same as the table attribute name.
    """
    ws = wb.create_sheet(title=table.element.get("name"))
    insert_table(table, ws, 1, 1)


@function_requires_deps("openpyxl", "premailer")
def document_to_workbook(doc, wb=None, base_url=None):
    """
    Takes a string representation of an html document and writes one sheet for
    every table in the document.
    The workbook is returned
    """
    if not wb:
        wb = Workbook()
        wb.remove(wb.active)

    inline_styles_doc = Premailer(
        doc, base_url=base_url, remove_classes=False
    ).transform()
    tables = get_Tables(inline_styles_doc)

    for table in tables:
        table_to_sheet(table, wb)

    return wb


def document_to_xl(doc, filename, base_url=None):
    """
    Takes a string representation of an html document and writes one sheet for
    every table in the document. The workbook is written out to a file called filename
    """
    wb = document_to_workbook(doc, base_url=base_url)
    wb.save(filename)


def insert_table(table, worksheet, column, row):
    """
    Inserts a table into the worksheet at the specified column and row
    """
    if table.head:
        row = write_rows(worksheet, table.head, row, column)
    if table.body:
        row = write_rows(worksheet, table.body, row, column)


def insert_table_at_cell(table, cell):
    """
    Inserts a table at the location of an openpyxl Cell object.
    """
    ws = cell.parent
    column, row = cell.column, cell.row
    insert_table(table, ws, column, row)
